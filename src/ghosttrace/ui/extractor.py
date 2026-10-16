import io
import os
import re

def extract_text_from_file(filename, file_bytes):
    """Extracts raw text from .docx, .pdf, .tex, .md, .txt, or .html files."""
    ext = os.path.splitext(filename)[1].lower()

    if ext in ['.txt', '.md', '.tex', '.html', '.htm']:
        return file_bytes.decode('utf-8', errors='replace')

    if ext in ['.docx', '.docs']:
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            full_text = []
            for para in doc.paragraphs:
                full_text.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells]
                    full_text.append(" | ".join(row_text))
            return "\n\n".join(full_text)
        except Exception as e:
            return f"[Error extracting text from Word document: {str(e)}]"

    if ext == '.pdf':
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            pages = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    pages.append(text)
            return "\n\n".join(pages)
        except Exception as e:
            return f"[Error extracting text from PDF: {str(e)}]"

    if ext == '.zip':
        try:
            import zipfile
            MAX_ZIP_MEMBER_SIZE = 10 * 1024 * 1024  # 10 MB per text file
            combined_texts = []
            with io.BytesIO(file_bytes) as buf, zipfile.ZipFile(buf) as z:
                # Priority: .tex files first, then .md, then .txt
                sorted_names = sorted(z.namelist(), key=lambda n: (not n.endswith('.tex'), not n.endswith('.md'), not n.endswith('.txt'), n))
                for name in sorted_names:
                    if name.startswith('__MACOSX') or name.endswith('/'):
                        continue
                    if name.endswith(('.tex', '.md', '.txt')):
                        info = z.getinfo(name)
                        if info.file_size > MAX_ZIP_MEMBER_SIZE:
                            combined_texts.append(f"% === File: {name} (Skipped: exceeds 10MB safety limit) ===")
                            continue
                        data = z.read(name).decode('utf-8', errors='replace')
                        combined_texts.append(f"% === File: {name} ===\n{data}")
            if combined_texts:
                return "\n\n".join(combined_texts)
            return "[No readable .tex, .md, or .txt files found in ZIP archive]"
        except Exception as e:
            return f"[Error extracting ZIP archive: {str(e)}]"

    # Fallback decode
    return file_bytes.decode('utf-8', errors='replace')
